"""Run export helpers."""

from __future__ import annotations

from .localization import tr

import csv
import json
from pathlib import Path
from html import escape

from .models import DesignCandidate, RunResult


CSV_HEADERS = [
    "Rank",
    "Trial",
    "Score_selection",
    "Defect",
    "ON_yield_pct",
    "Leak_pct",
    "dG_OFF_kcal_mol",
    "dG_ON_kcal_mol",
    "ddG_activation_kcal_mol",
    "dG_RBS_linker_kcal_mol",
    "Bad_RBS_linker",
    "STOP_codon",
    "First_STOP",
    "Reporter",
    "Trigger_DNA_5to3",
    "Switch_RNA_5to3",
    "Status",
    "Exclusion_reason",
]


def candidate_row(rank: int, candidate: DesignCandidate) -> list:
    return [
        "" if candidate.excluded else rank,
        candidate.trial,
        candidate.score,
        candidate.defect,
        candidate.on_yield_pct,
        candidate.leak_pct,
        candidate.delta_g_off,
        candidate.delta_g_on,
        candidate.ddg_activation,
        candidate.delta_g_rbs_linker,
        "" if candidate.has_bad_rbs_linker is None else ("yes" if candidate.has_bad_rbs_linker else "no"),
        "yes" if candidate.has_stop_codon else "no",
        candidate.first_stop or "",
        candidate.reporter_label,
        candidate.trigger_seq,
        candidate.switch_seq,
        "excluded" if candidate.excluded else "evaluated",
        tr(candidate.exclusion_reason),
    ]


def write_json(result: RunResult, output_dir: Path) -> Path:
    path = output_dir / f"{result.run_id}_results.json"
    path.write_text(json.dumps(result.to_dict(), indent=2), encoding="utf-8")
    return path


def write_csv(result: RunResult, output_dir: Path) -> Path:
    path = output_dir / f"{result.run_id}_ranked.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(CSV_HEADERS)
        for rank, candidate in enumerate(result.candidates, start=1):
            writer.writerow(candidate_row(rank, candidate))
    return path


def write_xlsx(result: RunResult, output_dir: Path) -> Path | None:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None

    path = output_dir / f"{result.run_id}_ranked.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = tr('Candidats AptaSwitch')
    sheet.append(CSV_HEADERS)
    for rank, candidate in enumerate(result.candidates, start=1):
        sheet.append(candidate_row(rank, candidate))

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    widths = [8, 8, 20, 10, 14, 12, 16, 16, 20, 20, 16, 12, 12, 14, 36, 100, 16, 54]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    workbook.save(path)
    return path


def write_top_svg(result: RunResult, output_dir: Path) -> Path | None:
    top = next((candidate for candidate in result.candidates if not candidate.excluded), None)
    if top is None:
        return None
    path = output_dir / f"{result.run_id}_top1.svg"
    trigger_width = max(180, min(900, len(top.trigger_seq) * 14))
    switch_width = max(260, min(1200, len(top.switch_seq) * 7))
    width = max(900, switch_width + 120)
    height = 380
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">
  <rect width="100%" height="100%" fill="#f8fafc"/>
  <text x="48" y="48" font-family="Arial" font-size="24" font-weight="700" fill="#111827">{escape(tr('Conception Apta2Switch-Studio la mieux classée'))}</text>
  <text x="48" y="78" font-family="Arial" font-size="14" fill="#374151">{escape(tr('Score {score:.4f} | ON {on:.2f}% | fuite {leak:.2f}% | ΔG RBS–linker {dg:.2f} kcal/mol', score=top.score, on=top.on_yield_pct, leak=top.leak_pct, dg=top.delta_g_rbs_linker))}</text>
  <rect x="48" y="118" width="{trigger_width}" height="34" rx="4" fill="#2563eb"/>
  <text x="60" y="141" font-family="Arial" font-size="13" font-weight="700" fill="#ffffff">{escape(tr('Trigger ADN 5′ vers 3′'))}</text>
  <text x="48" y="176" font-family="Menlo, Consolas, monospace" font-size="12" fill="#111827">{top.trigger_seq}</text>
  <rect x="48" y="218" width="{switch_width}" height="34" rx="4" fill="#059669"/>
  <text x="60" y="241" font-family="Arial" font-size="13" font-weight="700" fill="#ffffff">{escape(tr('Toehold switch ARN 5′ vers 3′'))}</text>
  <text x="48" y="276" font-family="Menlo, Consolas, monospace" font-size="11" fill="#111827">{top.switch_seq}</text>
  <text x="48" y="330" font-family="Arial" font-size="12" fill="#4b5563">{escape(tr('Ce SVG est un aperçu compact du résultat, pas une simulation de structure.'))}</text>
</svg>
"""
    path.write_text(svg, encoding="utf-8")
    return path


def write_run_log(result: RunResult, output_dir: Path) -> Path:
    path = output_dir / f"{result.run_id}.log"
    lines = [
        tr('Identifiant du run : {v0}', v0=result.run_id),
        tr('Statut : {v0}', v0=tr(result.status)),
        tr('Début : {v0}', v0=result.started_at),
        tr('Fin : {v0}', v0=result.completed_at),
        tr('Moteur : {v0}', v0=result.request.engine),
        tr('Essais demandés : {v0}', v0=result.request.trials),
        tr('Candidats : {v0}', v0=len(result.candidates)),
        tr('Exclure les candidats avec un STOP prématuré : {v0}', v0=result.request.exclude_stop_candidates),
        tr('Exclure les candidats au RBS–linker non linéaire : {v0}', v0=result.request.exclude_non_linear_candidates),
        tr('Candidats exclus (analyses restantes ignorées) : {v0}', v0=sum((candidate.excluded for candidate in result.candidates))),
        tr('Concentration du trigger (M) : {v0}', v0=result.request.thermo.resolved_trigger_concentration_m),
        tr('Concentration du toehold switch (M) : {v0}', v0=result.request.thermo.resolved_switch_concentration_m),
    ]
    if result.warnings:
        lines.append(tr('Avertissements :'))
        lines.extend(f"- {warning}" for warning in result.warnings)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_project_file(result: RunResult, output_dir: Path) -> Path:
    path = output_dir / ".aptaswitch.json"
    data = {
        "format": "aptaswitch-studio-project-v1",
        "run_id": result.run_id,
        "request": result.request.to_dict(),
        "results_file": f"{result.run_id}_results.json",
        "ranked_csv": f"{result.run_id}_ranked.csv",
        "ranked_xlsx": f"{result.run_id}_ranked.xlsx",
        "top_svg": f"{result.run_id}_top1.svg" if any(not candidate.excluded for candidate in result.candidates) else None,
    }
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return path


def export_run(result: RunResult) -> dict[str, str]:
    output_dir = Path(result.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {
        "json": str(write_json(result, output_dir)),
        "csv": str(write_csv(result, output_dir)),
        "log": str(write_run_log(result, output_dir)),
        "project": str(write_project_file(result, output_dir)),
    }
    xlsx = write_xlsx(result, output_dir)
    if xlsx:
        paths["xlsx"] = str(xlsx)
    svg = write_top_svg(result, output_dir)
    if svg:
        paths["svg"] = str(svg)
    return paths
