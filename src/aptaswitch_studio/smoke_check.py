"""Internal release check, executed by the packaged application itself."""

from __future__ import annotations

import json
import sys
import traceback
from io import BytesIO
from pathlib import Path
from unittest.mock import patch


def run_smoke_check(output_path: str) -> int:
    """Check shipped dependencies and every main page without running a design."""
    from aptaswitch_studio import __version__

    report: dict = {"ok": False, "frozen": bool(getattr(sys, "frozen", False)),
                    "version": __version__, "checks": []}
    try:
        import jinja2
        import numpy
        import pandas
        import scipy.interpolate
        import scipy.optimize
        import scipy.sparse
        import yaml
        from openpyxl import Workbook, load_workbook
        from packaging.tags import sys_tags
        from streamlit.testing.v1 import AppTest

        assert next(sys_tags()).interpreter == "cp311", "NUPACK wheels require the bundled Python 3.11"
        assert scipy.optimize.minimize_scalar(lambda x: (x - 2) ** 2).success
        assert scipy.sparse.eye(2).shape == (2, 2)
        assert yaml.safe_load("ready: true")["ready"]
        assert jinja2.Template("{{ ready }}").render(ready="ok") == "ok"
        assert pandas.DataFrame({"value": numpy.arange(2)}).shape == (2, 1)
        report["checks"].append("bundled scientific dependencies")

        from aptaswitch_studio.web_runtime import svg_to_png_bytes

        png = svg_to_png_bytes('<svg xmlns="http://www.w3.org/2000/svg" width="24" height="24">'
                               '<circle cx="12" cy="12" r="10" fill="#1258DC"/></svg>')
        assert png.startswith(b"\x89PNG\r\n\x1a\n")
        workbook = Workbook()
        workbook.active.append(["sequence", "ACGU"])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        assert load_workbook(stream).active["B1"].value == "ACGU"
        report["checks"].append("PNG and Excel exports")

        app_path = Path(__file__).with_name("web_app.py")
        assets = app_path.parent / "assets"
        for filename in ("logo.png", "app_icon.png", "igem_SU_logo.png"):
            assert (assets / filename).is_file(), filename
        # No dependency auto-detection, user settings or real calculations in CI.
        with (
            patch("aptaswitch_core.nupack_setup.load_settings", return_value={}),
            patch("aptaswitch_core.nupack_setup.save_settings"),
            patch("aptaswitch_core.nupack_setup.detect_nupack", return_value=None),
            patch("aptaswitch_core.nupack_setup.configured_nupack_import_path", return_value=None),
        ):
            app = AppTest.from_file(str(app_path), default_timeout=45)
            for language in ("en", "fr"):
                app.session_state["ui_language"] = language
                for page in ("Conception", "Résultats", "Réglages"):
                    app.session_state["nav"] = page
                    app.run()
                    failures = [str(error.value) for error in app.exception]
                    assert not failures, f"{language}/{page}: {failures}"
                    assert len(app.selectbox) + len(app.button) > 0
            report["checks"].append("Design, Results and Settings in English and French")
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    Path(output_path).write_text(json.dumps(report, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
