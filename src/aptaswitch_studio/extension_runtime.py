"""Background execution and reproducible exports for trigger extension searches."""

from __future__ import annotations

from aptaswitch_core.localization import get_language, language_context, tr

import csv
import json
import threading
import time
import zipfile
from contextvars import copy_context
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Mapping, Sequence

from aptaswitch_core.models import ProgressCallback, ThermoConditions
from aptaswitch_core.run_storage import new_run_directory
from aptaswitch_core.sequences import normalize_dna
from aptaswitch_core.extension_preview import ligand_positions
from aptaswitch_core.extension_visuals import export_extension_visuals
from aptaswitch_core.trigger_extension import (
    design_trigger_extensions,
    estimate_extension_search,
    extension_split,
    EXTENSION_SIDES,
)

from .web_runtime import JobSnapshot


@dataclass(frozen=True)
class ExtensionRunRequest:
    """Inputs to the NUPACK extension search at one or both ends of the core."""

    aptamer_dna: str
    trigger_dna: str
    target_lengths: tuple[int, ...]
    thermo: ThermoConditions
    import_path: str
    output_dir: Path
    molecule_name: str = "Système personnalisé"
    top: int = 100
    max_ensemble_candidates: int = 50000
    extension_side: str = "5prime"
    ligand_aptamer_positions: tuple[int, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        data = asdict(self)
        data["target_lengths"] = list(self.target_lengths)
        data["ligand_aptamer_positions"] = list(self.ligand_aptamer_positions)
        data["output_dir"] = str(self.output_dir)
        return data


_CANDIDATE_COLUMNS = (
    "rank", "rank_within_length", "final_score", "trigger_length", "extension_length",
    "extension_side", "extension_5prime", "extension_3prime", "extension_5prime_length", "extension_3prime_length",
    "extended_trigger", "mfe_structure", "mfe_energy_kcal_mol",
    "core_pair_probability_rmse", "core_pair_probability_mae",
    "reference_pair_probability_loss", "extension_pair_probability_sum",
    "extension_to_aptamer_probability_sum", "extension_to_trigger_core_probability_sum",
)


def _timestamped(message: str) -> str:
    return f"[{datetime.now().astimezone().strftime('%H:%M:%S')}] {message}"


def _candidate_table(candidates: Sequence[Mapping[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    """Keep all computed metrics, including nested probability data, in exports."""
    keys = set().union(*(candidate.keys() for candidate in candidates)) if candidates else set()
    columns = list(_CANDIDATE_COLUMNS)
    columns.extend(sorted(keys.difference(columns)))
    rows = [[_cell_value(candidate.get(column)) for column in columns] for candidate in candidates]
    return columns, rows


def _cell_value(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _write_xlsx(result: Mapping[str, Any], path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = openpyxl.Workbook()
    long_values: list[list[Any]] = []

    def append(sheet, values: Sequence[Any]) -> None:
        prepared = []
        for column, value in enumerate(values, start=1):
            value = _cell_value(value)
            if isinstance(value, str) and len(value) > 32767:
                # Excel silently truncates longer cells. Preserve every character
                # in numbered chunks instead, while the JSON also remains complete.
                address = f"{get_column_letter(column)}{sheet.max_row + 1}"
                parts = [value[start:start + 30000] for start in range(0, len(value), 30000)]
                long_values.extend([sheet.title, address, number, part] for number, part in enumerate(parts, start=1))
                value = tr('Voir Données longues · {v0}!{v1} · {v2} parties', v0=sheet.title, v1=address, v2=len(parts))
            prepared.append(value)
        sheet.append(prepared)

    candidates = workbook.active
    candidates.title = tr('Candidats étendus')
    append(candidates, headers)
    for row in rows:
        append(candidates, row)
    candidates.auto_filter.ref = candidates.dimensions
    for name, values in (
        (tr('Paramètres'), result.get("request", {})),
        (tr('Référence'), result.get("reference", {})),
        (tr('Recherche'), result.get("mfe_screen", {})),
    ):
        sheet = workbook.create_sheet(name)
        append(sheet, [tr('Champ'), tr('Valeur')])
        if isinstance(values, Mapping):
            for key, value in values.items():
                append(sheet, [key, value])
        else:
            append(sheet, [tr('Données'), values])
    curve_names = {
        "mfe_progress": tr('Progression MFE'),
        "ensemble_progress": tr('Progression ensemble'),
        "ranking": tr('Classement (courbe)'),
    }
    for key, records in result.get("curves", {}).items():
        sheet = workbook.create_sheet(curve_names.get(key, tr('Courbe')))
        if isinstance(records, list) and all(isinstance(record, Mapping) for record in records):
            columns = list(dict.fromkeys(field for record in records for field in record))
            if columns:
                append(sheet, columns)
                for record in records:
                    append(sheet, [record.get(column) for column in columns])
                sheet.auto_filter.ref = sheet.dimensions
            else:
                append(sheet, [tr('Aucune donnée')])
        else:
            append(sheet, [tr('Champ'), tr('Valeur')])
            append(sheet, [key, records])
    metadata = workbook.create_sheet(tr('Exécution'))
    append(metadata, [tr('Champ'), tr('Valeur')])
    for key in ("run_id", "status", "started_at", "completed_at", "warnings"):
        append(metadata, [key, result.get(key)])
    if long_values:
        sheet = workbook.create_sheet(tr('Données longues'))
        sheet.append([tr('Feuille'), tr('Cellule'), tr('Partie'), tr('Valeur')])
        for row in long_values:
            sheet.append(row)

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="193B35")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for index, column in enumerate(sheet.columns, start=1):
            length = max((len(str(cell.value or "")) for cell in column), default=12)
            sheet.column_dimensions[get_column_letter(index)].width = min(72, max(14, length + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    # Metadata is data, including arbitrary user-provided target names.
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _write_data_exports(result: Mapping[str, Any]) -> dict[str, str]:
    output_dir = Path(result["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    run_id = result["run_id"]
    json_path = output_dir / f"{run_id}_results.json"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    csv_path = output_dir / f"{run_id}_ranked.csv"
    headers, rows = _candidate_table(result.get("candidates", []))
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)
    xlsx_path = output_dir / f"{run_id}_ranked.xlsx"
    _write_xlsx(result, xlsx_path, headers, rows)
    paths = {"json": str(json_path), "csv": str(csv_path), "xlsx": str(xlsx_path)}
    paths.update(export_extension_visuals(result, output_dir))
    return paths


def _finish_exports(result: Mapping[str, Any], paths: dict[str, str], log: Sequence[str]) -> dict[str, str]:
    output_dir = Path(result["output_dir"])
    log_path = output_dir / f"{result['run_id']}.log"
    lines = [
        tr('Identifiant du run : {v0}', v0=result['run_id']),
        tr('Statut : {v0}', v0=tr(result['status'])),
        tr('Début : {v0}', v0=result['started_at']),
        tr('Fin : {v0}', v0=result['completed_at']),
        tr('Moteur : NUPACK — recherche exhaustive d’extension ({v0})', v0=tr(EXTENSION_SIDES[result.get('input', {}).get('extension_side', '5prime')])),
        tr('Candidats exportés : {v0}', v0=len(result.get('candidates', []))),
        "",
        *log,
    ]
    warnings = result.get("warnings", [])
    if warnings:
        lines.extend(["", tr('Avertissements :'), *(f"- {warning}" for warning in warnings)])
    log_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    paths["log"] = str(log_path)
    archive_path = output_dir / f"{result['run_id']}_complete.zip"
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path_value in paths.values():
            path = Path(path_value)
            archive.write(path, arcname=path.relative_to(output_dir))
    paths["zip"] = str(archive_path)
    return paths


def export_extension_result(result: Mapping[str, Any], *, log: Sequence[str] = ()) -> dict[str, str]:
    """Save all numerical results, visualizations, the execution log and a ZIP bundle."""
    return _finish_exports(result, _write_data_exports(result), log)


def run_extension_and_export(
    request: ExtensionRunRequest,
    *,
    progress_callback: ProgressCallback | None = None,
    cancel_check: Callable[[], bool] | None = None,
) -> tuple[dict[str, Any], dict[str, str]]:
    """Execute the scientific search and preserve its exact result and provenance."""
    started = datetime.now().astimezone()
    run_id, output_dir = new_run_directory(request.output_dir, "extension", started)
    log: list[str] = []
    current = 0
    trigger_length = len(normalize_dna(request.trigger_dna, field_name="trigger"))
    selected = ligand_positions(request.ligand_aptamer_positions, len(normalize_dna(request.aptamer_dna, field_name="aptamère")))
    estimate = estimate_extension_search(trigger_length, request.target_lengths)
    total = estimate["total_candidates_expected"]

    def progress(done: int, expected: int, message: str) -> None:
        nonlocal current, total
        current, total = done, expected
        log.append(_timestamped(message))
        if progress_callback:
            progress_callback(done, expected, message)

    progress(0, total, tr('Recherche NUPACK · {v0:,} extensions à examiner.', v0=total).replace(",", " "))
    result = dict(design_trigger_extensions(
        aptamer_dna=request.aptamer_dna,
        trigger_dna=request.trigger_dna,
        target_lengths=request.target_lengths,
        thermo=request.thermo,
        import_path=request.import_path,
        top=request.top,
        max_ensemble_candidates=request.max_ensemble_candidates,
        extension_side=request.extension_side,
        progress_callback=progress,
        cancel_check=cancel_check,
    ))
    result["input"]["ligand_aptamer_positions"] = list(selected)
    result.update({
        "run_id": run_id,
        "started_at": started.isoformat(timespec="seconds"),
        "completed_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "request": request.to_dict(),
        "output_dir": str(output_dir),
    })
    progress(current, total, tr('Écriture des séquences, métriques, structures 2D et courbes…'))
    exports = _write_data_exports(result)
    message = (
        tr('Extension annulée — résultats partiels exportés.')
        if result["status"] == "cancelled"
        else tr('Extension terminée — résultats exportés.')
    )
    progress(current, total, message)
    exports = _finish_exports(result, exports, log)
    return result, exports


class ExtensionRunJob:
    """A cancellable extension run with the same UI lifecycle as DesignRunJob."""

    def __init__(self, request: ExtensionRunRequest):
        extension_split(0, request.extension_side)
        self.request = request
        self._language = get_language()
        self._lock = threading.Lock()
        self._cancel = threading.Event()
        self._state = "running"
        self._current = 0
        trigger_length = len(normalize_dna(request.trigger_dna, field_name="trigger"))
        self._total = estimate_extension_search(trigger_length, request.target_lengths)["total_candidates_expected"]
        self._message = tr('Préparation de l’extension NUPACK…')
        self._started_monotonic = time.monotonic()
        self._phase_started_monotonic = self._started_monotonic
        self._phase_start_current = 0
        self._finished_monotonic: float | None = None
        self._log = [_timestamped(self._message)]
        self._result: dict[str, Any] | None = None
        self._exports: dict[str, str] | None = None
        self._error: str | None = None
        self._consumed = False
        # Preserve the language selected when this job was created.
        context = copy_context()
        self._thread = threading.Thread(target=context.run, args=(self._run,), daemon=True)

    def start(self) -> None:
        self._thread.start()

    def cancel(self) -> None:
        self._cancel.set()
        with self._lock, language_context(self._language):
            if self._state == "running":
                self._message = tr('Annulation demandée — fin du calcul NUPACK en cours…')
                self._append_log_locked(self._message)

    def _append_log_locked(self, message: str) -> None:
        self._log.append(_timestamped(message))
        # The complete log is retained in the exported journal.
        if len(self._log) > 500:
            self._log = self._log[-500:]

    def _progress(self, current: int, total: int, message: str) -> None:
        with self._lock:
            if current < self._current or total != self._total:
                self._phase_started_monotonic = time.monotonic()
                self._phase_start_current = current
            self._current = current
            self._total = total
            self._message = message
            self._append_log_locked(message)

    def _run(self) -> None:
        try:
            result, exports = run_extension_and_export(
                self.request,
                progress_callback=self._progress,
                cancel_check=self._cancel.is_set,
            )
        except Exception as exc:
            with self._lock:
                self._state = "failed"
                self._error = str(exc)
                self._message = tr('L’extension a échoué.')
                self._finished_monotonic = time.monotonic()
                self._append_log_locked(tr('ERREUR · {v0}', v0=exc))
            return
        with self._lock:
            self._state = "done"
            self._result = result
            self._exports = exports
            if result["status"] == "completed":
                self._current = self._total
            self._message = (
                tr('Extension annulée — résultats partiels exportés.')
                if result["status"] == "cancelled"
                else tr('Extension terminée.')
            )
            self._finished_monotonic = time.monotonic()
            self._append_log_locked(self._message)

    def snapshot(self) -> JobSnapshot:
        with self._lock:
            stopped_at = self._finished_monotonic or time.monotonic()
            elapsed = max(0.0, stopped_at - self._started_monotonic)
            eta = None
            phase_progress = self._current - self._phase_start_current
            if self._state == "running" and phase_progress > 0 and self._current < self._total:
                phase_elapsed = max(0.0, stopped_at - self._phase_started_monotonic)
                eta = max(0.0, phase_elapsed / phase_progress * (self._total - self._current))
            elif self._state != "running" or self._current >= self._total:
                eta = 0.0
            return JobSnapshot(
                state=self._state,
                current=self._current,
                total=self._total,
                message=self._message,
                result=self._result,
                exports=dict(self._exports) if self._exports else None,
                error=self._error,
                elapsed_seconds=elapsed,
                eta_seconds=eta,
                log=tuple(self._log),
            )

    def consume_once(self) -> bool:
        with self._lock:
            if self._consumed or self._state not in {"done", "failed"}:
                return False
            self._consumed = True
            return True
