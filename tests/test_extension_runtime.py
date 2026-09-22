from __future__ import annotations

import csv
import json
import sys
import tempfile
import threading
import unittest
import zipfile
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from aptaswitch_core.models import ThermoConditions
from aptaswitch_studio.extension_runtime import (
    ExtensionRunJob,
    ExtensionRunRequest,
    export_extension_result,
    run_extension_and_export,
)


def request(output_dir: Path, *, trigger: str = "ACGT") -> ExtensionRunRequest:
    return ExtensionRunRequest(
        aptamer_dna="ACGTACGT", trigger_dna=trigger, target_lengths=(5, 6),
        thermo=ThermoConditions(material="dna04", temperature_c=37),
        import_path="/configured/nupack", output_dir=output_dir,
        molecule_name="=custom system", top=12, max_ensemble_candidates=40,
    )


def core_result(status: str = "completed") -> dict:
    return {
        "status": status,
        "input": {"aptamer_dna": "ACGTACGT", "trigger_dna": "ACGT", "target_lengths": [5, 6]},
        "reference": {"mfe_structure": "........+....", "mfe_energy_kcal_mol": -1.2},
        "mfe_screen": {"total_candidates_expected": 20, "total_candidates_screened": 20},
        "candidates": [{
            "rank": 1, "rank_within_length": 1, "final_score": 0.25,
            "trigger_length": 5, "extension_length": 1, "extension_5prime": "T",
            "extended_trigger": "TACGT", "mfe_structure": "........+.....",
            "mfe_energy_kcal_mol": -1.21,
            "reference_pairs": [{"i": 0, "j": 9, "reference_probability": 0.8, "candidate_probability": 0.79}],
        }],
        "curves": {"reference_pairs": []},
        "warnings": [],
    }


def write_visual_fixture(result: dict, output_dir: Path) -> dict[str, str]:
    path = output_dir / f"{result['run_id']}_reference.svg"
    path.write_text('<svg xmlns="http://www.w3.org/2000/svg"/>', encoding="utf-8")
    return {"svg_reference": str(path)}


class ExtensionRuntimeTests(unittest.TestCase):
    def setUp(self):
        from aptaswitch_core.localization import language_context
        self.enterContext(language_context("fr"))

    def test_ligand_annotation_is_exported_without_changing_candidate_scores(self):
        with tempfile.TemporaryDirectory() as tmp:
            run_request = replace(request(Path(tmp)), ligand_aptamer_positions=(1, 4, 8))
            with (
                patch("aptaswitch_studio.extension_runtime.design_trigger_extensions", return_value=core_result()),
                patch("aptaswitch_studio.extension_runtime.export_extension_visuals", side_effect=write_visual_fixture),
            ):
                result, exports = run_extension_and_export(run_request)
            saved = json.loads(Path(exports["json"]).read_text())
            self.assertEqual(saved["input"]["ligand_aptamer_positions"], [1, 4, 8])
            self.assertEqual(saved["request"]["ligand_aptamer_positions"], [1, 4, 8])
            self.assertEqual(result["candidates"], core_result()["candidates"])

    def test_search_results_and_all_exports_retain_provenance_and_structures(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            run_request = request(Path(tmp))
            progress = []

            def search(**kwargs):
                kwargs["progress_callback"](12, 20, "Criblage · 12 extensions évaluées.")
                kwargs["progress_callback"](20, 20, "Validation thermodynamique terminée.")
                return core_result()

            with (
                patch("aptaswitch_studio.extension_runtime.design_trigger_extensions", side_effect=search) as engine,
                patch("aptaswitch_studio.extension_runtime.export_extension_visuals", side_effect=write_visual_fixture),
            ):
                result, exports = run_extension_and_export(
                    run_request, progress_callback=lambda *args: progress.append(args),
                )

            self.assertEqual(engine.call_count, 1)
            self.assertEqual(engine.call_args.kwargs["target_lengths"], (5, 6))
            self.assertEqual(engine.call_args.kwargs["thermo"], run_request.thermo)
            self.assertEqual(engine.call_args.kwargs["import_path"], "/configured/nupack")
            self.assertEqual(engine.call_args.kwargs["top"], 12)
            self.assertEqual(engine.call_args.kwargs["max_ensemble_candidates"], 40)
            self.assertEqual(set(exports), {"json", "csv", "xlsx", "log", "zip", "svg_reference"})
            self.assertEqual(Path(result["output_dir"]).parent, Path(tmp) / "extensions")
            self.assertRegex(result["run_id"], r"^\d{8}_\d{6}_extension_[a-f0-9]{6}$")
            saved = json.loads(Path(exports["json"]).read_text())
            self.assertEqual(saved["candidates"], core_result()["candidates"])
            self.assertEqual(saved["reference"], core_result()["reference"])
            self.assertEqual(saved["request"]["molecule_name"], "=custom system")
            self.assertEqual(saved["request"]["output_dir"], tmp)
            with Path(exports["csv"]).open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(rows[0]["extended_trigger"], "TACGT")
            self.assertEqual(rows[0]["mfe_structure"], "........+.....")
            self.assertEqual(json.loads(rows[0]["reference_pairs"]), core_result()["candidates"][0]["reference_pairs"])
            workbook = openpyxl.load_workbook(exports["xlsx"])
            self.assertIn("Référence", workbook.sheetnames)
            name_cells = [row[1] for row in workbook["Paramètres"] if row[0].value == "molecule_name"]
            self.assertEqual(name_cells[0].value, "=custom system")
            self.assertEqual(name_cells[0].data_type, "s")
            workbook.close()
            journal = Path(exports["log"]).read_text()
            self.assertRegex(journal, r"\[\d{2}:\d{2}:\d{2}\] Criblage · 12 extensions évaluées\.")
            self.assertIn("Extension terminée — résultats exportés.", journal)
            with zipfile.ZipFile(exports["zip"]) as archive:
                self.assertEqual(set(archive.namelist()), {Path(path).name for key, path in exports.items() if key != "zip"})
            self.assertTrue(any("structures 2D" in item[2] for item in progress))

    def test_progress_reports_elapsed_eta_and_normalizes_input_length(self):
        with patch("aptaswitch_studio.extension_runtime.time.monotonic", side_effect=[100.0, 110.0]):
            job = ExtensionRunJob(request(Path("/unused"), trigger="AC GT\n"))
            job._progress(5, 20, "Criblage MFE")
            snapshot = job.snapshot()
        self.assertEqual(snapshot.total, 20)
        self.assertEqual(snapshot.current, 5)
        self.assertEqual(snapshot.elapsed_seconds, 10)
        self.assertEqual(snapshot.eta_seconds, 30)
        self.assertRegex(snapshot.log[-1], r"^\[\d{2}:\d{2}:\d{2}\] Criblage MFE$")
        self.assertFalse(job.consume_once())

    def test_cancellation_reaches_engine_and_exports_partial_results(self):
        entered = threading.Event()
        released = threading.Event()
        with tempfile.TemporaryDirectory() as tmp:
            job = ExtensionRunJob(request(Path(tmp)))

            def search(**kwargs):
                entered.set()
                self.assertTrue(released.wait(timeout=5))
                self.assertTrue(kwargs["cancel_check"]())
                kwargs["progress_callback"](4, 20, "Annulation · 4 extensions évaluées.")
                result = core_result("cancelled")
                result["mfe_screen"]["total_candidates_screened"] = 4
                return result

            with (
                patch("aptaswitch_studio.extension_runtime.design_trigger_extensions", side_effect=search),
                patch("aptaswitch_studio.extension_runtime.export_extension_visuals", side_effect=write_visual_fixture),
            ):
                job.start()
                try:
                    self.assertTrue(entered.wait(timeout=5))
                    job.cancel()
                finally:
                    released.set()
                job._thread.join(timeout=5)
                self.assertFalse(job._thread.is_alive())
            snapshot = job.snapshot()
            self.assertEqual(snapshot.state, "done")
            self.assertEqual(snapshot.result["status"], "cancelled")
            self.assertEqual(snapshot.current, 4)
            self.assertEqual(snapshot.eta_seconds, 0)
            self.assertIn("partiels", snapshot.message)
            self.assertTrue(Path(snapshot.exports["zip"]).is_file())
            self.assertTrue(job.consume_once())
            self.assertFalse(job.consume_once())

    def test_eta_resets_when_search_moves_to_ensemble_validation(self):
        with patch("aptaswitch_studio.extension_runtime.time.monotonic", side_effect=[100.0, 120.0, 130.0]):
            job = ExtensionRunJob(request(Path("/unused")))
            job._progress(20, 20, "Criblage terminé")
            job._progress(0, 4, "Validation d’ensemble")
            job._progress(1, 4, "Validation d’ensemble 1/4")
            snapshot = job.snapshot()
        self.assertEqual(snapshot.elapsed_seconds, 30)
        self.assertEqual(snapshot.eta_seconds, 30)

    def test_engine_failure_is_reported_without_replacing_it_with_fake_results(self):
        job = ExtensionRunJob(request(Path("/unused")))
        with patch("aptaswitch_studio.extension_runtime.design_trigger_extensions", side_effect=RuntimeError("NUPACK indisponible")):
            job._run()
        snapshot = job.snapshot()
        self.assertEqual(snapshot.state, "failed")
        self.assertEqual(snapshot.error, "NUPACK indisponible")
        self.assertIsNone(snapshot.result)
        self.assertIsNone(snapshot.exports)
        self.assertTrue(any("ERREUR" in line for line in snapshot.log))
        self.assertTrue(job.consume_once())
        self.assertFalse(job.consume_once())

    def test_empty_cancelled_run_can_be_exported(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = core_result("cancelled")
            result.update({"run_id": "cancelled", "started_at": "now", "completed_at": "now", "output_dir": tmp})
            result["candidates"] = []
            result["reference"] = {}
            exports = export_extension_result(result)
            with Path(exports["csv"]).open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual(list(csv.DictReader(handle)), [])
            self.assertTrue(Path(exports["zip"]).is_file())

    def test_xlsx_preserves_long_curves_and_probability_data_without_cell_truncation(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as tmp:
            result = core_result()
            result.update({"run_id": "long_data", "started_at": "now", "completed_at": "now", "output_dir": tmp})
            points = [{"analyzed": index, "final_score": index / 1000} for index in range(1500)]
            result["curves"] = {"ensemble_progress": points}
            probabilities = [{"index": index, "probability": 0.123456789} for index in range(1500)]
            result["candidates"][0]["probabilities"] = probabilities
            with patch("aptaswitch_studio.extension_runtime.export_extension_visuals", return_value={}):
                exports = export_extension_result(result)
            workbook = openpyxl.load_workbook(exports["xlsx"])
            self.assertEqual(workbook["Progression ensemble"].max_row, 1501)
            self.assertEqual(workbook["Progression ensemble"].cell(1501, 1).value, 1499)
            chunks = [row[3].value for row in workbook["Données longues"].iter_rows(min_row=2)]
            self.assertEqual(json.loads("".join(chunks)), probabilities)
            self.assertTrue(all(len(chunk) <= 30000 for chunk in chunks))
            workbook.close()


if __name__ == "__main__":
    unittest.main()
